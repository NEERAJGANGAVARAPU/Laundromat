import os
import io
import csv
from datetime import datetime
from PIL import Image
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Sum, Q
from .models import Machine, LaundrySession
from .services import sync_expired_sessions

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def get_assets_dir():
    return os.path.join(os.path.dirname(__file__), 'assets')


def get_watermark_image_path():
    """
    Returns path to transparent watermark icon.
    Auto-generates 20% alpha transparent image if missing.
    """
    assets_dir = get_assets_dir()
    wm_path = os.path.join(assets_dir, 'sanielsoft_icon_watermark.png')
    if os.path.exists(wm_path):
        return wm_path

    icon_path = os.path.join(assets_dir, 'sanielsoft_icon.png')
    if not os.path.exists(icon_path):
        icon_path = os.path.abspath(
            os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'public', 'sanielsoft_icon.png')
        )

    if os.path.exists(icon_path):
        try:
            img = Image.open(icon_path).convert('RGBA')
            r, g, b, a = img.split()
            # 20% opacity (subtle watermark, 100% legible text on top)
            new_a = a.point(lambda p: int(p * 0.20))
            wm = Image.merge('RGBA', (r, g, b, new_a))
            os.makedirs(assets_dir, exist_ok=True)
            wm.save(wm_path)
            return wm_path
        except Exception:
            return icon_path
    return None


def get_full_logo_image_path():
    """
    Returns path to full SanielSoft logo image.
    """
    assets_dir = get_assets_dir()
    full_path = os.path.join(assets_dir, 'sanielsoft_full_logo.png')
    if os.path.exists(full_path):
        return full_path
    alt_path = os.path.abspath(
        os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'public', 'sanielsoft_full_logo.png')
    )
    if os.path.exists(alt_path):
        return alt_path
    return None


def draw_page_decorations(canvas, doc):
    """
    Renders on every page of the PDF:
    1. Background Watermark: Enlarged SanielSoft icon in center of page (semi-transparent)
    2. Footer: Divider line, Store label, Page number, and Powered by: SanielSoft full logo
    """
    canvas.saveState()

    # 1. Background Watermark: Enlarged SanielSoft Icon in center of page
    # Letter dimensions: 612pt wide x 792pt high.
    # Enlarged to ~50% page width (~300pt wide)
    watermark_path = get_watermark_image_path()
    if watermark_path and os.path.exists(watermark_path):
        wm_width = 300
        wm_height = 307
        wm_x = (612.0 - wm_width) / 2.0
        wm_y = (792.0 - wm_height) / 2.0
        canvas.drawImage(
            watermark_path,
            wm_x,
            wm_y,
            width=wm_width,
            height=wm_height,
            mask='auto',
            preserveAspectRatio=True
        )

    # 2. Bottom Footer on every page
    # Divider line above footer
    canvas.setStrokeColor(colors.HexColor('#E2E8F0'))
    canvas.setLineWidth(0.75)
    canvas.line(36, 44, 612 - 36, 44)

    # Left: Store & Report label
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#64748B'))
    canvas.drawString(36, 25, "Spinzone Laundromat · Official Session Usage Report")

    # Center: Page Number
    page_num = canvas.getPageNumber()
    canvas.drawCentredString(306, 25, f"Page {page_num}")

    # Right: "Powered by:" + SanielSoft full logo
    footer_logo_path = get_full_logo_image_path()
    if footer_logo_path and os.path.exists(footer_logo_path):
        logo_h = 16
        logo_w = 78
        logo_x = 612 - 36 - logo_w
        logo_y = 17
        canvas.drawRightString(logo_x - 6, 25, "Powered by:")
        canvas.drawImage(
            footer_logo_path,
            logo_x,
            logo_y,
            width=logo_w,
            height=logo_h,
            mask='auto',
            preserveAspectRatio=True
        )
    else:
        canvas.drawRightString(612 - 36, 25, "Powered by: SanielSoft")

    canvas.restoreState()



def format_dt(dt, fmt='%Y-%m-%d %H:%M:%S'):
    if not dt:
        return ''
    if timezone.is_aware(dt):
        dt = timezone.localtime(dt)
    return dt.strftime(fmt)

def get_day_wise_analytics(date_str=None, start_date=None, end_date=None):
    """
    Computes KPIs, machine utilization breakdown, and unique people count.
    Supports single date (date_str), date range (start_date to end_date), or all dates.
    """
    sync_expired_sessions()
    sessions = LaundrySession.objects.all().select_related('machine')

    # Available distinct dates for the date picker dropdown
    raw_dates = LaundrySession.objects.dates('started_at', 'day', order='DESC')
    available_dates = [d.strftime('%Y-%m-%d') for d in raw_dates]
    
    # Also ensure today's date is in available dates
    today_str = format_dt(timezone.now(), '%Y-%m-%d')
    if today_str not in available_dates:
        available_dates.insert(0, today_str)

    period_label = 'All Dates Combined'

    # Filter by date range or specific date if requested
    if start_date and end_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            e_d = datetime.strptime(end_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d, started_at__date__lte=e_d)
            period_label = f"{start_date} to {end_date}"
        except ValueError:
            pass
    elif start_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d)
            period_label = f"From {start_date}"
        except ValueError:
            pass
    elif date_str and date_str != 'all':
        try:
            selected_date = datetime.strptime(date_str, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date=selected_date)
            period_label = str(selected_date)
        except ValueError:
            pass

    total_sessions = sessions.count()
    completed_count = sessions.filter(status='completed').count()
    cancelled_count = sessions.filter(status='cancelled').count()
    active_count = sessions.filter(status='active').count()
    
    # Calculate unique customers and total minutes
    unique_customers = sessions.values('customer_email').distinct().count()
    total_minutes = sessions.aggregate(total=Sum('duration_minutes'))['total'] or 0

    # Machine-wise breakdown for the selected period / day
    # We include all active machines + any removed/inactive machines that have sessions!
    all_machines = Machine.objects.all().order_by('machine_code')
    machine_breakdown = []
    seen_codes = set()
    
    for m in all_machines:
        m_sessions = sessions.filter(Q(machine=m) | Q(machine_code=m.machine_code))
        runs_count = m_sessions.count()
        m_unique_people = m_sessions.values('customer_email').distinct().count()
        m_minutes = m_sessions.aggregate(total=Sum('duration_minutes'))['total'] or 0
        machine_status = m.status if m.is_active else 'removed'

        machine_breakdown.append({
            'code': m.machine_code,
            'type': m.machine_type,
            'runs_count': runs_count,
            'unique_people_count': m_unique_people,
            'total_minutes': m_minutes,
            'status': machine_status
        })
        seen_codes.add(m.machine_code)

    # In case there are historical sessions with a machine_code where the Machine object was purged
    orphaned_codes = sessions.exclude(machine_code__in=seen_codes).values_list('machine_code', flat=True).distinct()
    for code in orphaned_codes:
        if not code:
            continue
        m_sessions = sessions.filter(machine_code=code)
        runs_count = m_sessions.count()
        m_unique_people = m_sessions.values('customer_email').distinct().count()
        m_minutes = m_sessions.aggregate(total=Sum('duration_minutes'))['total'] or 0
        m_type = 'dryer' if code.upper().startswith('D') else 'washer'
        machine_breakdown.append({
            'code': code,
            'type': m_type,
            'runs_count': runs_count,
            'unique_people_count': m_unique_people,
            'total_minutes': m_minutes,
            'status': 'removed'
        })

    # Sort machines by runs_count descending
    machine_breakdown.sort(key=lambda x: x['runs_count'], reverse=True)
    most_active_machine = machine_breakdown[0]['code'] if machine_breakdown and machine_breakdown[0]['runs_count'] > 0 else 'N/A'

    return {
        'selected_date': period_label,
        'available_dates': available_dates,
        'summary': {
            'total_sessions': total_sessions,
            'unique_customers': unique_customers,
            'completed_count': completed_count,
            'cancelled_count': cancelled_count,
            'active_count': active_count,
            'total_minutes': total_minutes,
            'most_active_machine': most_active_machine,
        },
        'machine_breakdown': machine_breakdown
    }


def export_sessions_xls(date_str=None, start_date=None, end_date=None, status_filter=None, search_query=None):
    """
    Generates and streams a branded Excel (.xlsx) file of customer sessions and machine usage.
    Includes SanielSoft logo header, styled data table, and full logo footer.
    Supports single date, date range (start_date to end_date), or all dates.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    sync_expired_sessions()
    sessions = LaundrySession.objects.all().select_related('machine').order_by('-started_at')

    if start_date and end_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            e_d = datetime.strptime(end_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d, started_at__date__lte=e_d)
            period_label = f"{start_date} to {end_date}"
            filename_suffix = f"{start_date}_to_{end_date}"
        except ValueError:
            period_label = "Date Range"
            filename_suffix = "range"
    elif start_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d)
            period_label = f"From {start_date}"
            filename_suffix = f"from_{start_date}"
        except ValueError:
            period_label = "Date Range"
            filename_suffix = "range"
    elif date_str and date_str != 'all':
        try:
            d = datetime.strptime(date_str, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date=d)
            period_label = str(d)
            filename_suffix = str(d)
        except ValueError:
            period_label = str(date_str)
            filename_suffix = "date"
    else:
        period_label = "All Historical Records"
        filename_suffix = "all_time"

    if status_filter and status_filter != 'all':
        sessions = sessions.filter(status=status_filter)

    if search_query:
        from django.db.models import Q
        q = search_query.strip()
        sessions = sessions.filter(
            Q(customer_name__icontains=q) |
            Q(customer_email__icontains=q) |
            Q(customer_phone__icontains=q) |
            Q(machine__machine_code__icontains=q)
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Session Report"

    # Define styles
    brand_navy = '0B1A44'
    brand_blue = '2558EA'
    light_gray = 'F8FAFC'
    border_color = 'CBD5E1'
    text_dark = '1E293B'
    text_muted = '64748B'

    thin_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )

    title_font = Font(name='Calibri', size=16, bold=True, color=brand_navy)
    subtitle_font = Font(name='Calibri', size=11, bold=True, color=brand_blue)
    meta_font = Font(name='Calibri', size=9, color=text_muted)
    header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    data_font = Font(name='Calibri', size=9, color=text_dark)
    footer_font = Font(name='Calibri', size=9, italic=True, color=text_muted)

    header_fill = PatternFill(start_color=brand_navy, end_color=brand_navy, fill_type='solid')
    alt_row_fill = PatternFill(start_color=light_gray, end_color=light_gray, fill_type='solid')
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    # Set column widths
    col_widths = [12, 14, 14, 22, 30, 18, 16, 22, 22, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    current_row = 1

    # --- Logo Header ---
    full_logo_path = get_full_logo_image_path()
    if full_logo_path and os.path.exists(full_logo_path):
        try:
            logo_img = XlImage(full_logo_path)
            logo_img.width = 200
            logo_img.height = 41
            ws.add_image(logo_img, 'A1')
            current_row = 4  # Leave space for logo
        except Exception:
            current_row = 1

    # --- Title & Meta ---
    # "Spinzone" remains large (16pt); "| SanielSoft" is sized to match the subtitle below (11pt)
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    spinzone_inline_font = InlineFont(rFont='Calibri', sz=16.0, b=True, color=brand_navy)
    saniel_inline_font = InlineFont(rFont='Calibri', sz=11.0, b=True, color=brand_navy)
    title_rich = CellRichText(
        TextBlock(spinzone_inline_font, 'Spinzone '),
        TextBlock(saniel_inline_font, '| SanielSoft')
    )
    title_cell = ws.cell(row=current_row, column=1)
    title_cell.value = title_rich
    title_cell.alignment = left_align
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    sub_cell = ws.cell(row=current_row, column=1, value='Official Laundry Session & Machine Usage Report')
    sub_cell.font = subtitle_font
    sub_cell.alignment = left_align
    current_row += 1

    current_time_str = format_dt(timezone.now(), '%Y-%m-%d %I:%M %p')
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    meta_cell = ws.cell(row=current_row, column=1, value=f'Report Generated: {current_time_str}  |  Date Filter: {period_label}')
    meta_cell.font = meta_font
    meta_cell.alignment = left_align
    current_row += 2  # blank row spacer

    # --- Column Headers ---
    headers = [
        'Session ID', 'Machine Code', 'Machine Type', 'Customer Name',
        'Customer Email', 'Customer Phone', 'Duration (Min)',
        'Started At (Local)', 'Completed At (Local)', 'Status'
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1

    # --- Data Rows ---
    for row_idx, s in enumerate(sessions):
        local_started = format_dt(s.started_at)
        local_completed = format_dt(s.completed_at)
        m_code = s.machine_code or (s.machine.machine_code if s.machine else '')
        m_type = s.machine.machine_type if s.machine else ('dryer' if m_code.upper().startswith('D') else 'washer')
        phone_val = s.customer_phone if (s.customer_phone and s.customer_phone.strip() and s.customer_phone != 'Not provided') else (s.customer_phone or 'Not provided')

        row_data = [
            str(s.id),
            m_code,
            m_type,
            s.customer_name or 'Not specified',
            s.customer_email,
            phone_val,
            s.duration_minutes,
            local_started,
            local_completed,
            s.status.upper()
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = center_align if col_idx in (1, 2, 3, 7, 10) else left_align
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = alt_row_fill
        current_row += 1

    # --- Bottom Footer ---
    current_row += 1  # blank row spacer

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    footer1 = ws.cell(row=current_row, column=1, value='Powered by: SanielSoft (Saniel Software Services Pvt Ltd)')
    footer1.font = footer_font
    footer1.alignment = Alignment(horizontal='center', vertical='center')
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    footer2 = ws.cell(row=current_row, column=1, value='Spinzone Laundry Management Portal · All Rights Reserved')
    footer2.font = footer_font
    footer2.alignment = Alignment(horizontal='center', vertical='center')
    current_row += 2  # space before bottom logo

    # --- Full Logo at Bottom ---
    if full_logo_path and os.path.exists(full_logo_path):
        try:
            bottom_logo = XlImage(full_logo_path)
            bottom_logo.width = 200
            bottom_logo.height = 41
            # Center the logo roughly across the columns
            ws.add_image(bottom_logo, f'D{current_row}')
        except Exception:
            pass

    # Save to buffer and return response
    xls_buffer = io.BytesIO()
    wb.save(xls_buffer)
    xls_bytes = xls_buffer.getvalue()
    xls_buffer.close()

    response = HttpResponse(
        xls_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="laundromat_sessions_report_{filename_suffix}.xlsx"'
    return response


def export_sessions_pdf(date_str=None, start_date=None, end_date=None, status_filter=None, search_query=None):
    """
    Generates a structured, printable PDF usage report with KPI summary, machine utilization, and session logs.
    Supports single date, date range (start_date to end_date), or all dates.
    """
    analytics = get_day_wise_analytics(date_str=date_str, start_date=start_date, end_date=end_date)
    summary = analytics['summary']
    machine_breakdown = analytics['machine_breakdown']

    sessions = LaundrySession.objects.all().select_related('machine').order_by('-started_at')
    
    if start_date and end_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            e_d = datetime.strptime(end_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d, started_at__date__lte=e_d)
            period_label = f"{start_date} to {end_date}"
            filename_suffix = f"{start_date}_to_{end_date}"
        except ValueError:
            period_label = "Date Range"
            filename_suffix = "range"
    elif start_date:
        try:
            s_d = datetime.strptime(start_date, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date__gte=s_d)
            period_label = f"From {start_date}"
            filename_suffix = f"from_{start_date}"
        except ValueError:
            period_label = "Date Range"
            filename_suffix = "range"
    elif date_str and date_str != 'all':
        try:
            d = datetime.strptime(date_str, '%Y-%m-%d').date()
            sessions = sessions.filter(started_at__date=d)
            period_label = str(d)
            filename_suffix = str(d)
        except ValueError:
            period_label = str(date_str)
            filename_suffix = "date"
    else:
        period_label = "All Historical Records"
        filename_suffix = "all_time"

    if status_filter and status_filter != 'all':
        sessions = sessions.filter(status=status_filter)
    if search_query:
        from django.db.models import Q
        q = search_query.strip()
        sessions = sessions.filter(
            Q(customer_name__icontains=q) |
            Q(customer_email__icontains=q) |
            Q(customer_phone__icontains=q) |
            Q(machine__machine_code__icontains=q)
        )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0B1A44')
    )
    sub_style = ParagraphStyle(
        'RepSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#2558EA')
    )
    meta_style = ParagraphStyle(
        'RepMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748B')
    )
    table_text_style = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )
    table_head_style = ParagraphStyle(
        'TableHead',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    elements = []

    # Title & Subtitle (SanielSoft Branding)
    # "Spinzone" stays at title size; "| SanielSoft" matches subtitle size (11pt)
    elements.append(Paragraph("Spinzone <font size='11'>| SanielSoft</font>", title_style))
    elements.append(Paragraph("Official Laundry Session & Machine Usage Report", sub_style))
    
    current_time_str = format_dt(timezone.now(), '%Y-%m-%d %I:%M %p')
    elements.append(Paragraph(f"Report Generated: {current_time_str} | Date Filter: <b>{period_label}</b>", meta_style))
    elements.append(Spacer(1, 14))

    # KPI Summary Table
    elements.append(Paragraph("<b>1. Summary Metrics</b>", sub_style))
    elements.append(Spacer(1, 4))
    
    kpi_data = [
        [
            Paragraph("<b>Total Sessions:</b>", table_text_style), str(summary['total_sessions']),
            Paragraph("<b>Unique People:</b>", table_text_style), str(summary['unique_customers'])
        ],
        [
            Paragraph("<b>Completed Cycles:</b>", table_text_style), str(summary['completed_count']),
            Paragraph("<b>Cancelled Early:</b>", table_text_style), str(summary['cancelled_count'])
        ],
        [
            Paragraph("<b>Active Cycles:</b>", table_text_style), str(summary['active_count']),
            Paragraph("<b>Total Runtime:</b>", table_text_style), f"{summary['total_minutes']} mins"
        ],
        [
            Paragraph("<b>Most Active Machine:</b>", table_text_style), str(summary['most_active_machine']),
            "", ""
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[120, 150, 120, 150])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 14))

    # Machine Breakdown Table
    elements.append(Paragraph("<b>2. Machine Utilization & People Count</b>", sub_style))
    elements.append(Spacer(1, 4))

    mach_data = [[
        Paragraph("Machine Code", table_head_style),
        Paragraph("Type", table_head_style),
        Paragraph("Times Used", table_head_style),
        Paragraph("Unique Customers", table_head_style),
        Paragraph("Total Runtime", table_head_style),
        Paragraph("Current Status", table_head_style),
    ]]
    for m in machine_breakdown:
        mach_data.append([
            Paragraph(m['code'], table_text_style),
            Paragraph(m['type'].title(), table_text_style),
            Paragraph(str(m['runs_count']), table_text_style),
            Paragraph(str(m['unique_people_count']), table_text_style),
            Paragraph(f"{m['total_minutes']} mins", table_text_style),
            Paragraph(m['status'].title(), table_text_style),
        ])

    mach_table = Table(mach_data, colWidths=[80, 80, 85, 110, 95, 90])
    mach_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0B1A44')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(mach_table)
    elements.append(Spacer(1, 14))

    # Detailed Sessions Table
    elements.append(Paragraph("<b>3. Detailed Customer Sessions Log</b>", sub_style))
    elements.append(Spacer(1, 4))

    sess_data = [[
        Paragraph("Customer", table_head_style),
        Paragraph("Contact", table_head_style),
        Paragraph("Machine", table_head_style),
        Paragraph("Duration", table_head_style),
        Paragraph("Started (Local)", table_head_style),
        Paragraph("Status", table_head_style),
    ]]

    for s in sessions[:50]:  # Cap at top 50 rows for clean PDF layout
        started_str = format_dt(s.started_at, '%m/%d %I:%M %p') if s.started_at else '-'
        contact_str = f"{s.customer_email}<br/>{s.customer_phone or 'No phone'}"
        status_color = '#10B981' if s.status == 'completed' else '#EF4444' if s.status == 'cancelled' else '#2558EA'
        status_html = f"<font color='{status_color}'><b>{s.status.upper()}</b></font>"

        m_code = s.machine_code or (s.machine.machine_code if s.machine else '-')
        sess_data.append([
            Paragraph(s.customer_name or 'Walk-in Customer', table_text_style),
            Paragraph(contact_str, table_text_style),
            Paragraph(m_code, table_text_style),
            Paragraph(f"{s.duration_minutes} min", table_text_style),
            Paragraph(started_str, table_text_style),
            Paragraph(status_html, table_text_style),
        ])

    sess_table = Table(sess_data, colWidths=[105, 155, 60, 60, 95, 65])
    sess_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2558EA')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(sess_table)

    # Build PDF with background watermark and footer on every page
    doc.build(elements, onFirstPage=draw_page_decorations, onLaterPages=draw_page_decorations)
    pdf_bytes = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="laundromat_report_{filename_suffix}.pdf"'
    return response
